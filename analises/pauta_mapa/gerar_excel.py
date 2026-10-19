"""Planilha da pauta do Sacolão (Portaria SMSAN 175/2026) com preços oficiais da CEASA-PR.

Base: cotação oficial da CEASA-PR (DITEC), "preço mais comum" por produto, tipo e embalagem,
unidade Curitiba — tabela ceasa_pr_cotacoes, coletada da consulta oficial de evolução de
preços (mesmos valores do boletim "Coleta de Preços Diário"). Para cada espécie da pauta
entram TODOS os tipos oficiais do produto cotados na data mais recente (até 30 dias), e a
referência é a mediana do preço mais comum por kg — regra fixa, sem escolha manual.

Só lê o banco. Grava analises/pauta_mapa/pauta_sacolao_ceasa_<data>.xlsx.
    python analises/pauta_mapa/gerar_excel.py
"""

from __future__ import annotations

import os
import re
import sys
import unicodedata
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[1]
sys.path.insert(0, str(RAIZ))

from chat.pauta_dados import ATO, PAUTA_175  # noqa: E402

HOJE = date.today()
JANELA = 30
TETO_REVOGADO = 3.99

# espécie da pauta → nomes do produto na CEASA-PR (calendário oficial de comercialização e
# descrições da cotação de Curitiba); sem entrada = o próprio nome da espécie
NOMES_CEASA = {
    "alho-poró": ["alho porro", "alho poro"],
    "aipim": ["aipim mandioca", "mandioca"],
    "mandioquinha-salsa": ["mandioquinha batata salsa", "mandioquinha", "batata salsa"],
    "brócolis": ["couve brocolo", "couve brocolis", "brocolis"],
    "couve": ["couve manteiga"],
    "feijão-vagem": ["vagem", "feijao vagem"],
    "pitaya": ["pitaia", "pitaya"],
    "escarola": ["escarola chicoria", "escarola"],
    "maxixe": ["maxixi", "maxixe"],
    "inhame": ["inhame taia", "inhame"],
    "milho-verde": ["milho verde"],
    "cheiro-verde": ["cheiro verde"],
}
# Produtos que o boletim oficial lista separados (não são a espécie da pauta)
EXCLUIR = {"maçã": ["maca importada", "maca grand smith", "maca granny smith", "maca red delicious"],
           "pera": ["pera importada"], "batata": ["batata yakon"]}   # Granny Smith e Red Delicious
           # aparecem sob "MAÇÃ IMPORTADA" no boletim oficial "Coleta de Preços Diário"
# Art. 2º da Portaria 175: itens embalados e/ou minimamente processados ficam FORA da pauta
# (preço diferenciado) — saem do cálculo da referência
PROCESSADO = re.compile(r"\b(bj|pc)\b|bandeja|fatiad|descascad|picad|ralad|higienizad|\bbaby\b", re.I)
# Peso que a consulta de evolução não traz, mas o boletim oficial traz (Curitiba, 19/10/2026)
PESO_OFICIAL = {"MILHO VERDE SC C/50 ESPIGAS sc c/50 espigas": 13.0}

VARIEDADE = {"murcote": ["murcote", "murcot", "murkote", "morcote"], "tahiti": ["tahiti", "taiti"],
             "mimosa": ["mimosa"], "ponkan": ["ponkan", "poncam"], "pera": ["pera"], "lima": ["lima"],
             "rosa": ["rosa"]}

ARIAL = "Arial"
F_TIT = Font(name=ARIAL, bold=True, size=14)
F_CAB = Font(name=ARIAL, bold=True, color="FFFFFF")
F_TXT = Font(name=ARIAL, size=10)
F_INPUT = Font(name=ARIAL, size=10, color="0000FF")
F_LINK = Font(name=ARIAL, size=10, color="008000")
FILL_CAB = PatternFill("solid", fgColor="1F4E78")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_ALERTA = PatternFill("solid", fgColor="FCE4D6")
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))


def norm(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", s or "") if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[-/]", " ", s.lower())).strip()


def carregar(sb) -> list[dict]:
    desde = (HOJE - timedelta(days=JANELA)).isoformat()
    linhas, ultimo = [], 0
    while True:
        lote = (sb.table("ceasa_pr_cotacoes")
                .select("id,data_coleta,descricao,embalagem,peso_kg,preco,preco_kg")
                .eq("unidade", "CURITIBA").gte("data_coleta", desde).gt("id", ultimo)
                .order("id").limit(1000).execute().data or [])
        linhas += lote
        if len(lote) < 1000:
            return linhas
        ultimo = lote[-1]["id"]


def tipos_da_especie(cot: list[dict], especie: str, variedade: str | None, todas: list[str]) -> list[dict]:
    """Todos os tipos oficiais do produto, na data mais recente em que foi cotado."""
    nomes = [norm(n) for n in NOMES_CEASA.get(especie, [especie])]
    outros = [norm(n) for o in todas if o != especie for n in NOMES_CEASA.get(o, [o])]
    achados = []
    for c in cot:
        if not c["peso_kg"] and c["descricao"] in PESO_OFICIAL:
            c = {**c, "peso_kg": PESO_OFICIAL[c["descricao"]], "peso_do_boletim": True}
        d = norm(c["descricao"])
        casados = [n for n in nomes if d == n or d.startswith(n + " ")]
        if not casados:
            continue
        if any((d == o or d.startswith(o + " ")) and len(o) > max(map(len, casados)) for o in outros):
            continue                     # "batata doce ..." não entra em "batata"
        if any(d.startswith(norm(x)) for x in EXCLUIR.get(especie, [])):
            continue
        if PROCESSADO.search(c["descricao"] + " " + (c["embalagem"] or "")):
            continue
        if variedade and not any(v in d.split() or d.split()[1:2] == [v] or v in d
                                 for v in VARIEDADE.get(variedade, [norm(variedade)])):
            continue
        achados.append(c)
    if not achados:
        return []
    ultima = max(c["data_coleta"] for c in achados)
    return sorted([c for c in achados if c["data_coleta"] == ultima], key=lambda c: c["descricao"])


def cabecalho(ws, linha: int, titulos: list[str], larguras: list[int]):
    for i, (t, w) in enumerate(zip(titulos, larguras), start=1):
        cel = ws.cell(row=linha, column=i, value=t)
        cel.font, cel.fill = F_CAB, FILL_CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[linha].height = 32
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)
    ws.auto_filter.ref = f"A{linha}:{get_column_letter(len(titulos))}{linha}"


def main():
    from dotenv import load_dotenv
    from supabase import create_client
    load_dotenv(RAIZ / ".env")
    sb = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_KEY"])
    cot = carregar(sb)

    unicas = {}
    for l in PAUTA_175:
        unicas.setdefault((l["especie"], l["variedade"]), l)
    todas = sorted({e for e, _ in unicas})
    chave = lambda e, v: e + (f" {v}" if v else "")

    wb = Workbook()
    # ── Leia-me ──────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Leia-me"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    ws["A1"] = "Pauta do Sacolão × preços oficiais da CEASA-PR (Curitiba)"
    ws["A1"].font = F_TIT
    itens = [
        ("Gerado em", HOJE.strftime("%d/%m/%Y")),
        ("Pauta", f"{ATO['ato']}/{ATO['ano']} da SMSAN — Diário Oficial Eletrônico nº {ATO['diario_oficial']}, "
                  f"26/08/2026 (republicação do DOM nº 156, que saiu sem o Anexo I). {ATO['fonte']}"),
        ("Preços", "Cotação oficial da CEASA-PR (Divisão Técnica Econômica — DITEC), \"preço mais comum\" do "
                   "dia por produto, tipo e embalagem, unidade Curitiba. São os mesmos valores do boletim "
                   "\"Coleta de Preços Diário\" (https://www.ceasa.pr.gov.br/Pagina/Cotacao-Diaria-de-Precos-2026), "
                   "coletados pela plataforma na consulta oficial de evolução de preços."),
        ("Nomes dos produtos", "Calendário oficial de comercialização da CEASA-PR, unidade Curitiba "
                               "(cópia em docs/ceasa/calendario_comercializacao_ceasa-pr_curitiba.pdf): ex.: Couve "
                               "Manteiga, Brócolis, Tangerina Mimosa/Comum, Tangerina Murcot, Mandioca, Batata "
                               "Salsa. Ele confirma as leituras do Anexo I (mimosa e murcot são tangerinas; brócolis, "
                               "couve-flor e repolho estão entre as Hortaliças Folhosas). Boletim oficial de "
                               "19/10/2026 em docs/ceasa/coleta_precos_diario_curitiba_2026-10-19.pdf."),
        ("Regra do preço de referência", f"Para cada espécie da pauta entram TODOS os tipos oficiais do produto "
                                         f"cotados em Curitiba na data mais recente (até {JANELA} dias). Referência "
                                         "= mediana do preço mais comum por kg desses tipos; mínimo e máximo ao lado. "
                                         "Ficam fora: itens embalados ou minimamente processados (bandeja, pacote, "
                                         "fatiado, descascado, cenoura baby), que o Art. 2º da Portaria 175 trata fora "
                                         "da pauta; produtos que o boletim oficial lista à parte (Maçã Importada — inclui "
                                         "Granny Smith e Red Delicious —, Pera Importada, Batata Yakon); e tipos sem peso "
                                         "na embalagem (não têm R$/kg). Exceção: o milho-verde em saco de 50 espigas usa "
                                         "o peso de 13 kg informado no boletim oficial de 19/10/2026."),
        ("Teto de comparação (R$/kg)", TETO_REVOGADO),
        ("Atenção ao teto", "R$ 3,99/kg era o teto da Resolução SMSAN nº 2/2024, REVOGADA pela Resolução SMSAN nº "
                            "2/2026 (DOM 179, 24/09/2026). A Portaria 175 (Art. 1º, §1º) diz que o preço máximo é "
                            "definido por resolução específica, ainda não localizada. Edite a célula amarela para "
                            "testar outro teto: as abas recalculam."),
        ("Atacado × varejo", "Os preços da CEASA são de ATACADO; o teto vale para a venda ao consumidor no "
                             "Sacolão. A comparação indica pressão de custo, não descumprimento."),
        ("Abas", "Pauta 175 (86 linhas do Anexo I) · Referência por espécie (64 espécies, fórmulas) · "
                 "Cotações CEASA (todos os tipos usados, com R$/kg calculado)."),
        ("Cores", "Azul em fundo amarelo = valor editável · Verde = valor trazido de outra aba · Preto = fórmula."),
    ]
    for i, (k, v) in enumerate(itens, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(name=ARIAL, bold=True, size=10)
        c = ws.cell(row=i, column=2, value=v)
        c.font, c.alignment = F_TXT, Alignment(wrap_text=True, vertical="top")
    teto_linha = 3 + [k for k, _ in itens].index("Teto de comparação (R$/kg)")
    tc = ws.cell(row=teto_linha, column=2)
    tc.font, tc.fill, tc.number_format = F_INPUT, FILL_INPUT, '"R$" #,##0.00'
    tc.alignment = Alignment(horizontal="left")
    tc.comment = Comment("Resolução SMSAN 2/2024 (revogada em 24/09/2026). Edite para testar outro teto.", "AgroIA")
    TETO = f"'Leia-me'!$B${teto_linha}"

    # ── Cotações CEASA ───────────────────────────────────────────────────────
    wc = wb.create_sheet("Cotações CEASA")
    cabecalho(wc, 1, ["Espécie da pauta", "Descrição oficial (CEASA-PR)", "Embalagem", "Peso da embalagem (kg)",
                      "Preço mais comum (R$/embalagem)", "R$/kg", "Data da cotação"],
              [24, 52, 16, 14, 16, 12, 14])
    r = 2
    sem_peso = []
    for (esp, var), l in sorted(unicas.items(), key=lambda x: (x[1]["ordem_subgrupo"], x[0][0], x[0][1] or "")):
        for c in tipos_da_especie(cot, esp, var, todas):
            if not c["peso_kg"]:
                sem_peso.append((chave(esp, var), c["descricao"]))
                continue
            valores = [chave(esp, var), c["descricao"], c["embalagem"], float(c["peso_kg"]), float(c["preco"]),
                       f"=E{r}/D{r}", date.fromisoformat(c["data_coleta"])]
            for j, v in enumerate(valores, start=1):
                cel = wc.cell(row=r, column=j, value=v)
                cel.font = F_TXT
            wc.cell(row=r, column=4).number_format = "0.000"
            wc.cell(row=r, column=5).number_format = '"R$" #,##0.00'
            wc.cell(row=r, column=6).number_format = '"R$" #,##0.00'
            wc.cell(row=r, column=7).number_format = "dd/mm/yyyy"
            r += 1
    ult_cot = r - 1
    nota = r + 1
    wc.cell(row=nota, column=1, value="Fonte: CEASA-PR/DITEC, preço mais comum, unidade Curitiba; "
                                       "R$/kg = preço da embalagem ÷ peso da embalagem.").font = Font(name=ARIAL, italic=True, size=9)
    if sem_peso:
        wc.cell(row=nota + 1, column=1, value="Tipos sem peso na embalagem (fora do cálculo): " +
                "; ".join(f"{e}: {d}" for e, d in sem_peso)).font = Font(name=ARIAL, italic=True, size=9)

    # ── Referência por espécie ───────────────────────────────────────────────
    wr = wb.create_sheet("Referência por espécie", 1)
    cabecalho(wr, 1, ["Subgrupo", "Espécie da pauta", "Estações", "Tipos oficiais cotados",
                      "R$/kg referência (mediana)", "R$/kg mínimo", "R$/kg máximo", "Data da cotação",
                      "Acima do teto?", "Situação"],
              [32, 24, 12, 12, 16, 13, 13, 14, 12, 40])
    ESP, RKG, DT = (f"'Cotações CEASA'!$A$2:$A${ult_cot}", f"'Cotações CEASA'!$F$2:$F${ult_cot}",
                    f"'Cotações CEASA'!$G$2:$G${ult_cot}")
    estacoes = {}
    for l in PAUTA_175:
        estacoes.setdefault((l["especie"], l["variedade"]), set()).add(
            "Out/Inv" if l["estacao"] == "outono_inverno" else "Prim/Ver")
    r = 2
    for (esp, var), l in sorted(unicas.items(), key=lambda x: (x[1]["ordem_subgrupo"], x[0][0], x[0][1] or "")):
        k = chave(esp, var)
        wr.cell(row=r, column=1, value=l["subgrupo"])
        wr.cell(row=r, column=2, value=k)
        wr.cell(row=r, column=3, value=" + ".join(sorted(estacoes[(esp, var)])))
        wr.cell(row=r, column=4, value=f"=COUNTIFS({ESP},B{r})")
        wr[f"E{r}"] = ArrayFormula(f"E{r}", f"=IF(D{r}=0,\"\",MEDIAN(IF({ESP}=B{r},{RKG})))")
        wr.cell(row=r, column=6, value=f"=IF(D{r}=0,\"\",_xlfn.MINIFS({RKG},{ESP},B{r}))")
        wr.cell(row=r, column=7, value=f"=IF(D{r}=0,\"\",_xlfn.MAXIFS({RKG},{ESP},B{r}))")
        wr.cell(row=r, column=8, value=f"=IF(D{r}=0,\"\",_xlfn.MAXIFS({DT},{ESP},B{r}))")
        wr.cell(row=r, column=9, value=f"=IF(E{r}=\"\",\"\",IF(E{r}>{TETO},\"sim\",\"não\"))")
        wr.cell(row=r, column=10, value=f"=IF(D{r}=0,\"sem cotação em {JANELA} dias\",\"cotado\")")
        for j in range(1, 11):
            wr.cell(row=r, column=j).font = F_TXT
        for j in (5, 6, 7):
            wr.cell(row=r, column=j).number_format = '"R$" #,##0.00'
        wr.cell(row=r, column=8).number_format = "dd/mm/yyyy"
        r += 1
    ult_ref = r - 1
    r += 1
    resumo = [("Espécies com cotação", f"=COUNTIF(J2:J{ult_ref},\"cotado\")"),
              ("Espécies acima do teto", f"=COUNTIF(I2:I{ult_ref},\"sim\")"),
              ("Teto usado na comparação (R$/kg)", f"={TETO}")]
    for i, (k, f) in enumerate(resumo):
        wr.cell(row=r + i, column=2, value=k).font = Font(name=ARIAL, bold=True, size=10)
        c = wr.cell(row=r + i, column=5, value=f)
        c.font = F_LINK if "Leia-me" in f else F_TXT
    wr.cell(row=r + 2, column=5).number_format = '"R$" #,##0.00'

    # ── Pauta 175 ────────────────────────────────────────────────────────────
    wp = wb.create_sheet("Pauta 175", 1)
    cabecalho(wp, 1, ["Estação", "Grupo", "Subgrupo", "Mínimo do subgrupo", "Espécie", "Variedade",
                      "Texto no Anexo I", "R$/kg referência CEASA", "Acima do teto?", "Observação da transcrição"],
              [12, 12, 32, 10, 20, 12, 26, 14, 11, 60])
    REF_B, REF_E, REF_I = (f"'Referência por espécie'!$B$2:$B${ult_ref}",
                           f"'Referência por espécie'!$E$2:$E${ult_ref}",
                           f"'Referência por espécie'!$I$2:$I${ult_ref}")
    r = 2
    for l in PAUTA_175:
        vals = ["Outono/Inverno" if l["estacao"] == "outono_inverno" else "Primavera/Verão", l["grupo"],
                l["subgrupo"], l["minimo"], l["especie"], l["variedade"] or "", l["texto_original"],
                f"=IFERROR(INDEX({REF_E},MATCH(E{r}&IF(F{r}=\"\",\"\",\" \"&F{r}),{REF_B},0)),\"\")",
                f"=IFERROR(INDEX({REF_I},MATCH(E{r}&IF(F{r}=\"\",\"\",\" \"&F{r}),{REF_B},0)),\"\")",
                l["observacao"] or ""]
        for j, v in enumerate(vals, start=1):
            c = wp.cell(row=r, column=j, value=v)
            c.font = F_LINK if j in (8, 9) else F_TXT
            c.border = BORDA
        wp.cell(row=r, column=8).number_format = '"R$" #,##0.00'
        if l["validar_smsan"]:
            wp.cell(row=r, column=10).fill = FILL_ALERTA
        r += 1
    wp.cell(row=r + 1, column=1, value=f"Fonte: Anexo I da {ATO['ato']}/{ATO['ano']}, DOM nº {ATO['diario_oficial']} "
                                       "(26/08/2026). Mínimos por subgrupo e estação conforme o anexo; "
                                       "soma = 26 por estação.").font = Font(name=ARIAL, italic=True, size=9)

    destino = AQUI / f"pauta_sacolao_ceasa_{HOJE.isoformat()}.xlsx"
    wb.save(destino)
    print(destino, "| cotações:", ult_cot - 1, "| sem peso:", len(sem_peso))


if __name__ == "__main__":
    main()
